from pathlib import Path
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

source = Path(r"C:\Users\ebach\Desktop\10 inch Project\10-inch-FPV-drone-build-handover.docx")
doc = Document(source)

for item in doc.iter_inner_content():
    if isinstance(item, Paragraph):
        text = item.text.strip()
        if text:
            print(f"[{item.style.name}] {text}")
    elif isinstance(item, Table):
        print(f"[TABLE {len(item.rows)}x{len(item.columns)}]")
        for row in item.rows:
            cells = [" ".join(cell.text.split()) for cell in row.cells]
            print(" | ".join(cells))
        print("[/TABLE]")

for section_number, section in enumerate(doc.sections, start=1):
    header = " ".join(p.text.strip() for p in section.header.paragraphs if p.text.strip())
    footer = " ".join(p.text.strip() for p in section.footer.paragraphs if p.text.strip())
    if header:
        print(f"[HEADER {section_number}] {header}")
    if footer:
        print(f"[FOOTER {section_number}] {footer}")
