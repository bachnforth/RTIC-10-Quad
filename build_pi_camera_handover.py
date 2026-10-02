from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = r"C:\Users\ebach\Documents\ChatGPT\10 inch Quad\Pi-camera-integration-handover.docx"

BLACK = "000000"
NAVY = "1F4E78"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F2F2F2"
BORDER = "D9D9D9"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.find(qn("w:tcMar"))
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn("w:" + tag))
        if node is None:
            node = OxmlElement("w:" + tag)
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def set_repeat_heading_style(style):
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.keep_together = True


def apply_font(run, name="Aptos", size=10.5, bold=False, color=BLACK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.alignment = 1
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, text in enumerate(headers):
        cell = hdr.cells[idx]
        cell.text = text
        set_cell_shading(cell, NAVY)
        set_cell_borders(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            apply_font(run, size=9.5, bold=True, color=WHITE)
        if widths:
            cell.width = widths[idx]
    for r_idx, row_data in enumerate(rows):
        row = table.add_row()
        for c_idx, text in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.text = str(text)
            if r_idx % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            set_cell_borders(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    apply_font(run, size=9.3)
            if widths:
                cell.width = widths[c_idx]
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.add_run(text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    apply_font(run, name="Cascadia Mono", size=8.8)
    return p


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.68)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.08

title_style = styles["Title"]
title_style.font.name = "Aptos Display"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
title_style.font.size = Pt(24)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_style.paragraph_format.space_after = Pt(5)
title_ppr = title_style._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

for style_name, size, before, after in (
    ("Heading 1", 15, 12, 5),
    ("Heading 2", 12, 8, 3),
):
    st = styles[style_name]
    st.font.name = "Aptos Display"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    set_repeat_heading_style(st)

for list_name in ("List Bullet", "List Bullet 2", "List Number"):
    styles[list_name].font.name = "Aptos"
    styles[list_name].font.size = Pt(10.2)
    styles[list_name].paragraph_format.space_after = Pt(2.5)

# Title block
p = doc.add_paragraph(style="Title")
p.add_run("Raspberry Pi Camera Integration Handover")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(10)
r = p.add_run("10 inch FPV drone payload and analog video path")
apply_font(r, size=12, bold=False, color="404040")

meta = add_table(
    doc,
    ["Field", "Current record"],
    [
        ("Prepared", "22 September 2026"),
        ("Aircraft project", "Rekon10 10 inch quadcopter on the existing 4S system"),
        ("Payload", "Raspberry Pi Zero 2 W and Raspberry Pi Camera Module 2"),
        ("Purpose", "Continue camera, analog-video, and power integration without losing verified settings or safety gates"),
    ],
    [Inches(1.55), Inches(5.35)],
)

p = doc.add_paragraph()
r = p.add_run("Current conclusion. ")
apply_font(r, bold=True)
p.add_run(
    "The Pi and camera work on the bench, and the Pi is configured to generate an NTSC composite preview automatically. "
    "The remaining gates are to finish and inspect the Pi-to-flight-controller wiring, choose a power source with adequate margin, "
    "and prove the complete image path in the goggles under a sustained propeller-off load test. The system is not yet cleared for flight."
)

doc.add_heading("Status at a glance", level=1)
add_table(
    doc,
    ["Item", "Status", "Evidence or remaining work"],
    [
        ("Pi operating system", "Verified", "Raspberry Pi OS Lite 64-bit based on Debian 13 Trixie"),
        ("Camera hardware", "Verified", "IMX219 Camera Module 2 detected; 3280 x 2464 still capture passed"),
        ("Remote access", "Verified previously", "SSH key login worked at the last known address; the Pi was offline when this handover was generated"),
        ("Composite output", "Configured", "NTSC output enabled and a continuous 640 x 480, 30 fps DRM preview service installed"),
        ("Pi video pads", "User reported complete", "Wires soldered to Pi TV and GND pads; solder quality and continuity not independently inspected"),
        ("FC and VTX", "Partially complete", "VTX is reported soldered to the LUX F765; Pi wire terminations at FC Camera/GND are not verified"),
        ("Pi power", "Open decision", "Matek MBEC6S candidate is only 1.5A continuous; FC 5V rail is shared and requires load-margin validation"),
        ("Goggle image and latency", "Not tested", "No confirmed live image, OSD overlay, dropout test, or end-to-end latency measurement"),
        ("Flight readiness", "Not cleared", "Power, RF, video, and hover revalidation remain incomplete"),
    ],
    [Inches(1.48), Inches(1.25), Inches(4.17)],
)

doc.add_heading("Verified Pi configuration", level=1)
add_table(
    doc,
    ["Setting", "Value"],
    [
        ("Hostname", "quadpi"),
        ("Username", "ebach"),
        ("Last known IPv4 address", "192.168.40.193, assigned by DHCP"),
        ("Wi-Fi MAC", "88:A2:9E:C4:B8:13"),
        ("Wi-Fi", "RoboticsLan on 2.4 GHz; the Windows computer may use RoboticsLan-5G on the bridged subnet"),
        ("Camera", "Sony IMX219, Camera Module 2"),
        ("Last successful camera test", "3280 x 2464 JPEG; no corruption observed"),
        ("Last successful power check", "vcgencmd get_throttled returned 0x0"),
        ("Camera software", "rpicam-apps-lite; rpicam-hello available"),
    ],
    [Inches(2.05), Inches(4.85)],
)

doc.add_heading("Composite preview software", level=2)
p = doc.add_paragraph()
p.add_run("Boot configuration changes").bold = True
add_bullet(doc, "In /boot/firmware/config.txt: dtoverlay=vc4-kms-v3d,composite")
add_bullet(doc, "In /boot/firmware/cmdline.txt: vc4.tv_norm=NTSC")
add_bullet(doc, "Original files preserved as config.txt.pre-fpv and cmdline.txt.pre-fpv")

p = doc.add_paragraph()
p.add_run("Automatic preview service").bold = True
add_code(doc, "quadpi-camera-preview.service")
add_code(doc, "/usr/bin/rpicam-hello --timeout 0 --fullscreen --viewfinder-width 640 --viewfinder-height 480 --framerate 30")
p = doc.add_paragraph(
    "The service was enabled, reboot-tested, and observed active. Its journal reported a DRM preview window and the 640 x 480 IMX219 stream."
)

doc.add_heading("Required wiring", level=1)
p = doc.add_paragraph(
    "The Pi camera is converted to analog composite video by the Pi. The flight controller then inserts its analog OSD before sending video to the Rush Tank Solo VTX."
)
add_table(
    doc,
    ["Connection", "From", "To", "State"],
    [
        ("Camera data", "Camera Module 2", "Pi CSI connector through the 15-pin-to-22-pin ribbon", "Connected and tested"),
        ("Composite signal", "Pi TV test pad", "LUX F765 pad labelled Camera", "Pi end soldered; FC end unverified"),
        ("Video reference", "Pi GND", "LUX F765 GND", "Pi end soldered; FC end unverified"),
        ("Processed video", "LUX F765 VTX pad", "Rush Tank Solo video input", "User reports VTX already soldered"),
        ("VTX RF", "Rush Tank Solo MMCX", "Rated 5.8 GHz antenna", "Must be attached before VTX power"),
        ("Optional control", "LUX TX4 S/A", "Tank Solo SmartAudio", "Not required to obtain an image"),
    ],
    [Inches(1.25), Inches(1.62), Inches(2.72), Inches(1.31)],
)

p = doc.add_paragraph()
r = p.add_run("Video path. ")
apply_font(r, bold=True)
p.add_run("Camera Module 2 -> CSI ribbon -> Pi -> TV composite output -> FC Camera input -> FC OSD -> FC VTX output -> Tank Solo -> goggles")

doc.add_heading("Power decision and risk", level=1)
p = doc.add_paragraph()
r = p.add_run("Preferred design. ")
apply_font(r, bold=True)
p.add_run(
    "Use a dedicated 4S-input regulator providing a measured 5.0 to 5.1V at 3A continuous or greater. "
    "This isolates the flight-video computer from the flight controller's accessory rail and preserves current margin."
)

p = doc.add_paragraph()
r = p.add_run("Rejected candidate. ")
apply_font(r, bold=True)
p.add_run(
    "The photographed Matek MBEC6S has a selectable 5V or 9V output but is rated for only 1.5A continuous. "
    "It may boot the Pi under light load, but it does not meet the planned 3A continuous margin and should not be used for flight-critical video."
)

p = doc.add_paragraph()
r = p.add_run("Proposed FC-powered alternative. ")
apply_font(r, bold=True)
p.add_run(
    "The boxed LUX F765 pads labelled 5V and GND are electrically the correct pair for Pi power. "
    "However, the LUX manual rates the entire onboard 5V regulator at 2.5A, shared with the flight controller and other 5V equipment. "
    "This configuration couples a Pi overload or short to the flight-control power system and must not be accepted without a sustained load test."
)

doc.add_heading("Next safe milestone", level=1)
p = doc.add_paragraph(
    "Complete a propeller-off bench validation of the entire camera power and video chain. Do not proceed directly to a hover."
)

doc.add_heading("1 Inspect and meter with all power removed", level=2)
for text in (
    "Remove every propeller and disconnect the flight battery, USB cables, and CanaKit supply.",
    "Photograph both sides of the Pi and FC solder joints under good light and magnification.",
    "Confirm the Pi TV wire terminates only at the LUX Camera pad, not Camera Ctrl, VTX, a UART, or a power pad.",
    "Confirm the Pi ground terminates at LUX ground and that no loose strand can contact an adjacent pad.",
    "Check for an unintended short between Pi TV and GND and between Pi 5V and GND.",
    "Verify that the VTX antenna is attached before any VTX power is applied.",
):
    add_number(doc, text)

doc.add_heading("2 Validate the selected power source", level=2)
for text in (
    "Power the regulator or FC without the Pi connected and measure polarity at the intended Pi leads.",
    "Confirm 5.0 to 5.1V. Stop immediately for 9V, battery voltage, reversed polarity, unstable readings, or unexpected heating.",
    "Connect the Pi only after the unloaded voltage and polarity pass.",
    "Never connect the CanaKit adapter and airborne 5V source simultaneously.",
):
    add_number(doc, text)

doc.add_heading("3 Run the sustained propeller off test", level=2)
for text in (
    "Use a smoke stopper for the first low-load aircraft power-up. Keep the VTX at its minimum practical bench power.",
    "Verify that the FC, receiver, Pi, camera service, VTX, and goggles all start without resets.",
    "Measure voltage at the Pi while the camera preview is active. Treat voltage near or below 4.8V as inadequate margin.",
    "Run the complete system for at least 15 minutes while watching the goggle image and checking component temperature.",
    "Check vcgencmd get_throttled before and after the test. Any nonzero undervoltage history requires correction before flight.",
    "Move wiring gently and exercise the radio controls while watching for image breakup, Pi reboot, receiver dropout, or FC reset.",
):
    add_number(doc, text)

doc.add_heading("4 Verify the pilot image", level=2)
for text in (
    "Confirm correct camera orientation, usable field of view, stable NTSC sync, and visible FC OSD.",
    "Measure practical camera-to-goggle latency and compare it with the course's accepted setup before using it as the pilot view.",
    "Confirm that the image recovers automatically after a normal power cycle.",
    "Record the VTX channel, power, antenna, measured Pi voltage, test duration, temperatures, and any dropouts.",
):
    add_number(doc, text)

doc.add_heading("Stop conditions", level=1)
add_table(
    doc,
    ["Stop immediately if", "Required response"],
    [
        ("BEC or FC output is not a stable 5V", "Disconnect the Pi and correct the power source before continuing"),
        ("Pi reports undervoltage or reboots", "Increase power margin, shorten or enlarge power wiring, and repeat the full test"),
        ("FC or receiver resets when the Pi starts", "Do not use the shared FC rail; install a dedicated regulator"),
        ("VTX has no antenna attached", "Remove power and attach the correct antenna before retrying"),
        ("Image rolls, flickers, disappears, or has severe motor-noise sensitivity", "Inspect grounds and routing; do not fly from the feed"),
        ("Any solder bridge, exposed conductor, smoke, odor, or rapid heating is observed", "Remove power immediately and inspect under magnification"),
    ],
    [Inches(3.05), Inches(3.85)],
)

doc.add_heading("Remote access and service commands", level=1)
p = doc.add_paragraph(
    "The last known address is DHCP-assigned and may change. Check the router or scan for MAC 88:A2:9E:C4:B8:13 if 192.168.40.193 stops responding."
)
add_code(doc, "ssh -i C:\\Users\\ebach\\.ssh\\id_ed25519_rpi ebach@192.168.40.193")
add_table(
    doc,
    ["Purpose", "Command on the Pi"],
    [
        ("Preview service status", "systemctl status quadpi-camera-preview.service"),
        ("Preview service log", "journalctl -u quadpi-camera-preview.service -b --no-pager"),
        ("Restart preview", "sudo systemctl restart quadpi-camera-preview.service"),
        ("List detected cameras", "rpicam-hello --list-cameras"),
        ("Check undervoltage history", "vcgencmd get_throttled"),
        ("Check temperature", "vcgencmd measure_temp"),
        ("Clean shutdown", "sudo shutdown -h now"),
    ],
    [Inches(2.25), Inches(4.65)],
)

doc.add_heading("Known maintenance item", level=1)
p = doc.add_paragraph(
    "A full operating-system upgrade was attempted but stopped safely during package download because the Raspberry Pi repository slowed to a multi-hour estimate. "
    "No package installation had begun, and the package database was clean afterward. Complete updates later on reliable power and network service, then retest the camera service after reboot."
)

doc.add_heading("Records to preserve", level=1)
for text in (
    "Clear photographs of Pi TV, GND, 5V, and power-ground joints and the matching FC terminations.",
    "Power-source make, model, current rating, output selection, measured unloaded voltage, and loaded voltage at the Pi.",
    "FC 5V rail loads if the shared rail is used.",
    "Goggle model, receiver module, VTX channel and power, antenna type, image orientation, and observed latency.",
    "Fifteen-minute test results, temperatures, get_throttled value, and any reset or dropout.",
):
    add_bullet(doc, text)

doc.add_heading("References", level=1)
refs = (
    "Raspberry Pi Zero 2 W test pads: https://datasheets.raspberrypi.com/rpizero2/raspberry-pi-zero-2-w-test-pads.pdf",
    "Raspberry Pi composite output configuration: https://www.raspberrypi.com/documentation/computers/config_txt.html",
    "Raspberry Pi camera software: https://www.raspberrypi.com/documentation/computers/camera_software.html",
    "Lumenier LUX F765 NDAA manual: https://getfpv-media.s3.us-east-1.amazonaws.com/wysiwyg/files/Lumenier_LUX_F765_FC_NDAA_User_Manual_v1.pdf",
    "Matek MBEC6S product page: https://www.mateksys.com/?portfolio=mbec6s",
)
for ref in refs:
    add_bullet(doc, ref)

# Footer with page field
footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = fp.add_run("Raspberry Pi camera integration handover  |  ")
apply_font(run, size=8.5, color="666666")
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
fp._p.append(fld)

# Prevent table rows from splitting across pages where practical.
for table in doc.tables:
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)

doc.core_properties.title = "Raspberry Pi Camera Integration Handover"
doc.core_properties.subject = "10 inch FPV drone Pi Zero 2 W and Camera Module 2 integration"
doc.core_properties.author = "Project handover"
doc.core_properties.keywords = "Raspberry Pi Zero 2 W, Camera Module 2, LUX F765, composite video, FPV"
doc.save(OUTPUT)
print(OUTPUT)
